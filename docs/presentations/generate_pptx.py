"""
StudyOS – Profesyonel Sunum Uretici (Turkce)
Cikti: docs/presentations/StudyOS-Presentation-v1.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "StudyOS-Presentation-v1.pptx")

# ── Renk Paleti ───────────────────────────────────────────────────────────────
BG_DARK   = RGBColor(0x0D, 0x17, 0x2B)
BG_CARD   = RGBColor(0x16, 0x25, 0x40)
BG_DARK2  = RGBColor(0x05, 0x0E, 0x1C)
ACCENT    = RGBColor(0x38, 0xBD, 0xF8)   # elektrik mavi
ACCENT2   = RGBColor(0x0E, 0xA5, 0xE9)
GREEN     = RGBColor(0x34, 0xD3, 0x99)   # zumrut yesil
AMBER     = RGBColor(0xFB, 0xBF, 0x24)   # amber
PURPLE    = RGBColor(0xC0, 0x84, 0xFC)   # mor
RED       = RGBColor(0xF8, 0x71, 0x71)   # kirmizi
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREY      = RGBColor(0x94, 0xA3, 0xB8)
DARKBLUE  = RGBColor(0x0E, 0x2A, 0x44)
DARKGREEN = RGBColor(0x0B, 0x2A, 0x1A)
DARKNAVY  = RGBColor(0x12, 0x1E, 0x33)
FONT      = "Segoe UI"

# Slayt boyutu 16:9
W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

TOTAL = 19


# ── Yardimci Fonksiyonlar ─────────────────────────────────────────────────────

def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    return s


def rect(s, x, y, w, h, fill=None, line=None, lw=Pt(0)):
    sh = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    ln = sh.line
    if line:
        ln.color.rgb = line
        ln.width = lw
    else:
        ln.fill.background()
    return sh


def rrect(s, x, y, w, h, fill=BG_CARD):
    sh = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def bar(s, x, y, w, thick=0.045, color=ACCENT):
    sh = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(thick))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def dot(s, x, y, d=0.24, color=ACCENT):
    sh = s.shapes.add_shape(9, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def tb(s, text, x, y, w, h,
       size=24, bold=False, color=WHITE,
       align=PP_ALIGN.LEFT, italic=False):
    txb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.name   = FONT
    f.size   = Pt(size)
    f.bold   = bold
    f.italic = italic
    f.color.rgb = color
    return txb


def bullets(s, items, x, y, w, h, size=16, color=WHITE, sub_color=GREY, bullet="->"):
    txb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, str):
            text, level = item, 0
        else:
            text, level = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(3 if level == 0 else 1)
        r = p.add_run()
        pre = (bullet + "  ") if (level == 0 and bullet) else ("    -  " if level == 1 else "")
        r.text = pre + text
        f = r.font
        f.name  = FONT
        f.size  = Pt(size if level == 0 else size - 2)
        f.bold  = False
        f.color.rgb = color if level == 0 else sub_color
    return txb


def tag(s, label, x, y, w=2.0, h=0.42, bg=ACCENT2, fg=WHITE, size=13):
    sh = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = bg
    sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    f = r.font
    f.name  = FONT
    f.size  = Pt(size)
    f.bold  = True
    f.color.rgb = fg
    return sh


def snum(s, n):
    tb(s, f"{n:02d} / {TOTAL}", 12.3, 7.1, 0.9, 0.3,
       size=11, color=GREY, align=PP_ALIGN.RIGHT)


def section(s, label, x=0.4, y=0.15):
    tb(s, label.upper(), x, y, 8, 0.32,
       size=11, color=ACCENT, bold=True)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 01 — KAPAK
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG_DARK2

bar(s, 3.5, 3.42, 6.33, thick=0.055, color=ACCENT)

tb(s, "StudyOS", 0, 1.7, 13.33, 1.4,
   size=88, bold=True, align=PP_ALIGN.CENTER)

tb(s, "Yapay Zeka Destekli Egitim Ekosistemi",
   0, 3.62, 13.33, 0.7,
   size=26, color=GREY, align=PP_ALIGN.CENTER)

tb(s, "Ogrenci Mobil Uygulamasi   .   Kurumsal Yonetim Platformu",
   0, 4.5, 13.33, 0.5,
   size=18, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

tb(s, "Iki urun.  Tek ekosistem.  Tek hedef.",
   0, 5.15, 13.33, 0.5,
   size=16, color=GREY, align=PP_ALIGN.CENTER)

snum(s, 1)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 02 — SORUN: OGRENCILER
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Sorun")
bar(s, 0.4, 0.55, 1.2)

tb(s, "Ogrenciler dagitik araclarla calisiyor.",
   0.4, 0.65, 12.0, 0.85,
   size=40, bold=True)

cards = [
    ("Dagitik\nAraclar"),
    ("Sinav\nAnalizi Yok"),
    ("Geri Bildirim\nYok"),
    ("Tutarlilik\nTakibi Yok"),
    ("Sessiz\nBasarisizlik"),
]
icons = ["Planlayici", "Sonuc Defteri", "Hata Dongusu", "Pomodoro", "Uyari"]
cx = 0.4
for i, label in enumerate(cards):
    rrect(s, cx, 1.72, 2.3, 2.25, fill=BG_CARD)
    bar(s, cx, 1.72, 2.3, thick=0.06, color=ACCENT)
    tb(s, label, cx, 2.2, 2.3, 1.3,
       size=17, color=WHITE, align=PP_ALIGN.CENTER)
    cx += 2.47

rect(s, 0.4, 4.15, 12.5, 0.85, fill=DARKBLUE)
tb(s, "YKS  .  LGS  .  KPSS  —  Milyonlarca ogrenci.  Hicbir birlesik arac.",
   0.6, 4.25, 12.1, 0.6,
   size=20, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

tb(s, "Plan yok.  Takip yok.  Rehberlik yok.  Kurum baglantisi yok.",
   0.4, 5.2, 12.5, 0.5,
   size=16, color=GREY, align=PP_ALIGN.CENTER)

snum(s, 2)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 03 — SORUN: KURUMLAR
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Sorun")
bar(s, 0.4, 0.55, 1.2)

tb(s, "Kurumlar kör bir sekilde yonetiyor.",
   0.4, 0.65, 12.0, 0.85,
   size=40, bold=True)

rows = [
    ("Yoklama",          "Kagit formlar ve Excel tablolari"),
    ("Odev Takibi",      "WhatsApp mesajlari ve sozlu talimatlar"),
    ("Sinav Sonuclari",  "Optik cevap kagitlarindan elle girilen veriler"),
    ("Veli Iletisimi",   "Gayri resmi arama ve mesajlar"),
    ("Riskli Ogrenci",   "Erken uyari mekanizmasi bulunmuyor"),
]
ry = 1.72
for title, sub in rows:
    rrect(s, 0.4, ry, 12.5, 0.72, fill=BG_CARD)
    bar(s, 0.4, ry, 12.5, thick=0.045, color=DARKBLUE)
    tb(s, title, 0.6, ry + 0.1, 4.0, 0.5, size=18, bold=True)
    tb(s, sub,   4.8, ry + 0.12, 7.8, 0.48, size=16, color=GREY)
    ry += 0.82

tb(s, "Kurum ogrencilerinin varligini biliyor.  Nasil olduklarini bilmiyor.",
   0.4, 6.1, 12.5, 0.5,
   size=17, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

snum(s, 3)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 04 — NEDEN ONEMLI
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Onemin Boyu")
bar(s, 0.4, 0.55, 1.2)

rect(s, 0.4, 0.9, 12.5, 2.3, fill=DARKBLUE)
bar(s, 0.4, 0.9, 0.1, thick=2.3, color=ACCENT)
tb(s, '"En onemli veri, onu en cok\nihtiyac duyan insanlara gorunmez."',
   0.75, 1.0, 12.0, 2.1,
   size=38, bold=True, italic=True, align=PP_ALIGN.CENTER)

bar(s, 2.5, 3.35, 8.33, color=ACCENT)

for cx, heading, items, col in [
    (0.5, "Ogrenciler Icin", [
        "Olcemedigini gelistiremezsin",
        "Dagitik araclar tutarsiligi bozuyor",
        "Rehberlik olmadan emek bosa gidiyor",
    ], ACCENT),
    (7.0, "Kurumlar Icin", [
        "Erken uyari yoksa mudahale gecikir",
        "Manuel surecler olceklenmez",
        "Gayri resmi iletisim veli guvenini erozyor",
    ], GREEN),
]:
    tb(s, heading, cx, 3.55, 5.8, 0.5, size=20, bold=True, color=col)
    for i, item in enumerate(items):
        tb(s, "->  " + item, cx, 4.15 + i * 0.68, 5.8, 0.6, size=17)

snum(s, 4)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 05 — MEVCUT COZUMLER VE SINIRLILIKLAR
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Pazar Boslugu")
bar(s, 0.4, 0.55, 1.2)

tb(s, "Mevcut araclar tek sorun cozer.\nStudyOS sistemi cozer.",
   0.4, 0.65, 12.5, 1.25, size=36, bold=True)

headers = ["Arac Turu", "Ne Yapar", "Neyi Kacirir"]
col_w   = [3.2, 3.8, 5.1]
col_x   = [0.4, 3.65, 7.5]

for i, (h, w, x) in enumerate(zip(headers, col_w, col_x)):
    rect(s, x, 2.1, w - 0.05, 0.45, fill=ACCENT2)
    tb(s, h, x + 0.1, 2.15, w - 0.2, 0.38, size=15, bold=True)

table_rows = [
    ("Genel calisma uygulamalari", "Zamanlayici - Notlar",        "AI yok - Kurum baglantisi yok"),
    ("Okul yonetim yazilimlari",  "Yoklama - Program",           "Ogrenci uygulamasi yok - Analitik yok"),
    ("Mesajlasma uygulamalari",   "Iletisim",                    "Egitim icin tasarlanmamis - Veri yok"),
    ("Excel tablolari",           "Manuel takip",                "Zeka yok - Hata orani yuksek"),
    ("Ayri LMS platformlari",     "Icerik dagitimi",             "Kisisel calisma verisiyle baglanmaz"),
]
alts = [BG_CARD, DARKNAVY]
for ri, (t, does, miss) in enumerate(table_rows):
    ry = 2.6 + ri * 0.65
    bg = alts[ri % 2]
    for x, w in zip(col_x, col_w):
        rect(s, x, ry, w - 0.05, 0.62, fill=bg)
    tb(s, t,    col_x[0] + 0.1, ry + 0.07, col_w[0] - 0.2, 0.5, size=14)
    tb(s, does, col_x[1] + 0.1, ry + 0.07, col_w[1] - 0.2, 0.5, size=14, color=GREY)
    tb(s, miss, col_x[2] + 0.1, ry + 0.07, col_w[2] - 0.2, 0.5, size=14, color=AMBER)

tb(s, "->  Hicbir arac ogrencinin kisisel calisma verisini kurum yonetimiyle baglamadi.  Ta ki simdi.",
   0.4, 5.98, 12.5, 0.55,
   size=16, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

snum(s, 5)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 06 — STUDYOS'U TANIYALIM
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()

tb(s, "Tanitiyoruz:", 0, 0.55, 13.33, 0.65,
   size=22, color=GREY, align=PP_ALIGN.CENTER)
tb(s, "StudyOS", 0, 1.05, 13.33, 1.3,
   size=80, bold=True, align=PP_ALIGN.CENTER)

bar(s, 4.2, 2.52, 4.93, color=ACCENT)

# Sol kart - Ogrenci
rrect(s, 0.5, 2.8, 5.4, 3.5, fill=BG_CARD)
bar(s, 0.5, 2.8, 5.4, thick=0.07, color=ACCENT)
tb(s, "Ogrenci Mobil Uygulamasi",
   0.5, 3.05, 5.4, 0.62,
   size=21, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
tb(s, "Kisisellestirilmis yapay zeka\ndestekli akademik rehber",
   0.5, 3.75, 5.4, 0.85,
   size=17, color=GREY, align=PP_ALIGN.CENTER)
tb(s, "Android  .  iOS",
   0.5, 5.35, 5.4, 0.5,
   size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Orta baglanti
rect(s, 5.92, 3.55, 1.49, 1.6, fill=DARKBLUE)
tb(s, "<->", 5.92, 3.7, 1.49, 0.7,
   size=28, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
tb(s, "Ortak\nBulut", 5.92, 4.35, 1.49, 0.72,
   size=13, color=GREY, align=PP_ALIGN.CENTER)

# Sag kart - Kurum
rrect(s, 7.43, 2.8, 5.4, 3.5, fill=BG_CARD)
bar(s, 7.43, 2.8, 5.4, thick=0.07, color=GREEN)
tb(s, "Kurumsal Web Platformu",
   7.43, 3.05, 5.4, 0.62,
   size=21, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
tb(s, "Her kurumun ihtiyacina ozel\nbütünlesik yonetim katmani",
   7.43, 3.75, 5.4, 0.85,
   size=17, color=GREY, align=PP_ALIGN.CENTER)
tb(s, "Web  .  Duyarli Tasarim",
   7.43, 5.35, 5.4, 0.5,
   size=14, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

tb(s, "Iki urun.  Tek ekosistem.  Tek hedef.",
   0, 6.5, 13.33, 0.55,
   size=18, color=GREY, italic=True, align=PP_ALIGN.CENTER)

snum(s, 6)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 07 — URUN VIZYONU
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Vizyon")
bar(s, 0.4, 0.55, 1.2)

rect(s, 0.4, 0.9, 12.5, 3.4, fill=DARKBLUE)
bar(s, 0.4, 0.9, 0.1, thick=3.4, color=ACCENT)
tb(s,
   '"StudyOS; her ogrencinin ogrenme yolculugunu yonetebildigi\n'
   've her egitim kurumunun ogrenci tabanini tam anlamiyla\n'
   'yonetebildigi tek, bagli, yapay zeka destekli\n'
   'egitim ekosistemi haline gelecektir."',
   0.75, 1.05, 12.0, 3.1,
   size=27, italic=True, align=PP_ALIGN.CENTER)

tb(s, "Misyon:", 0.4, 4.58, 2.2, 0.45, size=16, bold=True, color=ACCENT)
tb(s, "Calisma ile basari arasindaki surtuculugu ortadan kaldirmak.",
   2.7, 4.58, 10.2, 0.45, size=18, bold=True)

commitments = [
    "Her ogrenci kisisel ve akilli bir akademik rehbere kavusur",
    "Her kurum tam operasyonel gorunum elde eder",
    "Iki taraf her zaman otomatik olarak bagli kalir",
]
for i, c in enumerate(commitments):
    tb(s, "->  " + c, 0.4, 5.25 + i * 0.52, 12.5, 0.5, size=17, color=GREY)

snum(s, 7)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 08 — OGRENCI PLATFORMU
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Ogrenci Platformu")
bar(s, 0.4, 0.55, 2.2, color=ACCENT)

tb(s, "Ogrenci Mobil Uygulamasi",
   0.4, 0.65, 12.5, 0.85, size=38, bold=True)
tb(s, "Tek uygulama.  Tam aciklik.",
   0.4, 1.45, 12.5, 0.5, size=20, color=GREY)

cols = [
    ("PLANLA",    ACCENT,  ["Gunluk calisma plani", "Sinav / Universite modu"]),
    ("TAKIP ET",  GREEN,   ["Konu ve soru takibi", "Yanlis defteri", "Akilli tekrar sistemi"]),
    ("ANALIZ ET", AMBER,   ["Deneme sinavi analizi", "Istatistik ve grafikler", "AI tarafindan olusturulan plan"]),
    ("ODAKLAN",   PURPLE,  ["Pomodoro zamanlayici", "Akilli bildirimler", "Bulut senkronizasyonu"]),
]
cx = 0.4
for heading, color, items in cols:
    rrect(s, cx, 2.1, 3.1, 4.5, fill=BG_CARD)
    bar(s, cx, 2.1, 3.1, thick=0.07, color=color)
    tb(s, heading, cx, 2.22, 3.1, 0.52,
       size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
    iy = 2.88
    for item in items:
        tb(s, ".  " + item, cx + 0.15, iy, 2.8, 0.55, size=16)
        iy += 0.56
    cx += 3.23

rect(s, 0.4, 6.78, 12.5, 0.55, fill=DARKBLUE)
tb(s, "AI Calisma Kocu — Kisisellestirilmis akademik rehberlik, her an erisim",
   0.6, 6.83, 12.1, 0.45,
   size=17, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

snum(s, 8)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 09 — KURUMSAL PLATFORM
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Kurumsal Platform")
bar(s, 0.4, 0.55, 2.4, color=GREEN)

tb(s, "Kurumsal Web Platformu",
   0.4, 0.65, 12.5, 0.85, size=38, bold=True)
tb(s, "Tam kontrol.  Tam gorunum.",
   0.4, 1.45, 12.5, 0.5, size=20, color=GREY)

quads = [
    ("YONET",    ACCENT, ["Kurum ve sube kaydı", "Sinif - Ogretmen - Ogrenci", "Rol tabanli yetkiler"]),
    ("ILET",     GREEN,  ["Dijital yoklama", "PDF ve video icerik", "Odevler ve duyurular"]),
    ("DEGERLEN", AMBER,  ["Deneme sinavi olusturma", "Optik kagit aktarimi", "Konu bazli basari analizi"]),
    ("ANLA",     PURPLE, ["AI akademik danısman", "Riskli ogrenci tespiti", "Disari aktarilabilir raporlar"]),
]
cx = 0.4
for heading, color, items in quads:
    rrect(s, cx, 2.1, 3.1, 4.5, fill=BG_CARD)
    bar(s, cx, 2.1, 3.1, thick=0.07, color=color)
    tb(s, heading, cx, 2.22, 3.1, 0.52,
       size=15, bold=True, color=color, align=PP_ALIGN.CENTER)
    iy = 2.88
    for item in items:
        tb(s, ".  " + item, cx + 0.15, iy, 2.8, 0.55, size=15)
        iy += 0.57
    cx += 3.23

rect(s, 0.4, 6.78, 12.5, 0.55, fill=DARKBLUE)
tb(s, "Veli Erisimi — Aile gorunurlugu icin ozel, guvenilir, yapisal kanal",
   0.6, 6.83, 12.1, 0.45,
   size=17, bold=True, color=GREEN, align=PP_ALIGN.CENTER)

snum(s, 9)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 10 — YAPAY ZEKA OZELLIKLERI
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Yapay Zeka")
bar(s, 0.4, 0.55, 1.5)

tb(s, "AI bir ozellik degil — omurgadir.",
   0.4, 0.65, 12.5, 0.85, size=40, bold=True)

# Ogrenci AI karti
rrect(s, 0.4, 1.75, 5.7, 4.85, fill=BG_CARD)
bar(s, 0.4, 1.75, 5.7, thick=0.07, color=ACCENT)
tb(s, "Ogrenci AI", 0.55, 1.88, 5.4, 0.55,
   size=20, bold=True, color=ACCENT)
ai_s = [
    "AI Calisma Kocu — kisisel rehberlik",
    "AI Plan Olusturucu — uyarlanabilir program",
    "Akilli Tekrar — zayif konulari onceliklendir",
    "Ilerleme analizi ve istatistikler",
]
for i, item in enumerate(ai_s):
    tb(s, "->  " + item, 0.55, 2.56 + i * 0.78, 5.35, 0.7, size=17)

# Ortak rozet
rrect(s, 5.85, 3.68, 1.63, 1.0, fill=ACCENT2)
tb(s, "<->\nPaylasili", 5.85, 3.73, 1.63, 0.9,
   size=15, bold=True, align=PP_ALIGN.CENTER)

# Kurum AI karti
rrect(s, 7.23, 1.75, 5.7, 4.85, fill=BG_CARD)
bar(s, 7.23, 1.75, 5.7, thick=0.07, color=GREEN)
tb(s, "Kurum AI", 7.38, 1.88, 5.4, 0.55,
   size=20, bold=True, color=GREEN)
ai_k = [
    "AI Akademik Danısman — kurumsal oneriler",
    "Riskli Ogrenci Tespiti — erken uyari",
    "Konu bazli basari analizi",
    "Ogrenci gelisim raporlari",
]
for i, item in enumerate(ai_k):
    tb(s, "->  " + item, 7.38, 2.56 + i * 0.78, 5.35, 0.7, size=17)

rect(s, 0.4, 6.78, 12.5, 0.55, fill=DARKBLUE)
tb(s, "Her iki platformda olusturulan AI analizleri her iki tarafta da aninda erisilebilinir.",
   0.6, 6.83, 12.1, 0.45,
   size=17, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

snum(s, 10)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 11 — UYGULAMALAR NASIL BIRLIKTE CALISIYOR
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Entegrasyon")
bar(s, 0.4, 0.55, 1.5)

tb(s, "Tek bagli ekosistem.",
   0.4, 0.65, 12.5, 0.85, size=40, bold=True)

# Ogrenci dugumu
rrect(s, 0.5, 1.72, 3.5, 2.0, fill=BG_CARD)
bar(s, 0.5, 1.72, 3.5, thick=0.06, color=ACCENT)
tb(s, "Ogrenci Uygulamasi", 0.65, 1.85, 3.2, 0.5, size=16, bold=True, color=ACCENT)
tb(s, "Android . iOS", 0.65, 2.32, 3.2, 0.38, size=13, color=GREY)
tb(s, "Kisisel calisma verisi\ngercek zamanli islenir",
   0.65, 2.67, 3.2, 0.72, size=13)

# Bulut dugumu
rect(s, 4.92, 2.02, 3.5, 1.62, fill=DARKBLUE)
bar(s, 4.92, 2.02, 3.5, thick=0.06, color=ACCENT2)
tb(s, "StudyOS Bulut", 5.07, 2.15, 3.2, 0.48, size=16, bold=True, color=ACCENT)
tb(s, "Hesaplar . API\nVeritabani . AI",
   5.07, 2.65, 3.2, 0.72, size=13, color=GREY)

# Kurum dugumu
rrect(s, 9.23, 1.72, 3.5, 2.0, fill=BG_CARD)
bar(s, 9.23, 1.72, 3.5, thick=0.06, color=GREEN)
tb(s, "Kurumsal Platform", 9.38, 1.85, 3.2, 0.5, size=16, bold=True, color=GREEN)
tb(s, "Web . Duyarli", 9.38, 2.32, 3.2, 0.38, size=13, color=GREY)
tb(s, "Yonetim verisi\ngercek zamanli islenir",
   9.38, 2.67, 3.2, 0.72, size=13)

# Oklar
tb(s, "<-------------------->", 3.6, 2.62, 2.1, 0.5,
   size=16, color=ACCENT, align=PP_ALIGN.CENTER)
tb(s, "<-------------------->", 8.1, 2.62, 2.1, 0.5,
   size=16, color=ACCENT, align=PP_ALIGN.CENTER)

# Veri akisi tablosu
hdr_texts = [
    "Ogrenci Uygulamasindan  ->  Kuruma",
    "Kurumdan  ->  Ogrenci Uygulamasina",
]
col2_x = [0.4, 6.97]
col2_w = [6.3, 6.3]
ry = 4.05
for hx, hw, ht in zip(col2_x, col2_w, hdr_texts):
    rect(s, hx, ry, hw - 0.08, 0.45, fill=ACCENT2)
    tb(s, ht, hx + 0.1, ry + 0.05, hw - 0.2, 0.38, size=14, bold=True)

lefts  = ["Calisma sure kayitlari", "Deneme sinavi sonuclari", "AI performans verisi", "Basari grafikleri"]
rights = ["Ogretmen tarafindan verilen odevler", "PDF dokuman ve materyaller", "AI analizleri ve raporlar", "Duyurular"]
for i, (l, r) in enumerate(zip(lefts, rights)):
    bg = BG_CARD if i % 2 == 0 else DARKNAVY
    ry2 = 4.55 + i * 0.5
    rect(s, 0.4,  ry2, 6.22, 0.47, fill=bg)
    rect(s, 6.97, ry2, 6.22, 0.47, fill=bg)
    tb(s, "->  " + l, 0.55, ry2 + 0.06, 6.0, 0.38, size=14)
    tb(s, "->  " + r, 7.1, ry2 + 0.06, 6.0, 0.38, size=14)

snum(s, 11)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 12 — KULLANICI YOLCULUGU
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Kullanici Yolculugu")
bar(s, 0.4, 0.55, 1.8)

tb(s, "Ayse ile tanisalim — 17 yasinda, YKS'ye hazirlaniyor.",
   0.4, 0.65, 12.5, 0.75, size=33, bold=True)
tb(s, "Dershanesi StudyOS kullaniyor.",
   0.4, 1.3, 12.5, 0.45, size=20, color=GREY)

steps = [
    ("Pazartesi\nSabah",   "Uygulama haftalik plani\nzayif konulara gore hazirlar", ACCENT),
    ("Pazartesi\nAksam",   "AI kocu atlanan\nGeometri'yi isaret eder",              ACCENT),
    ("Carsamba",           "Ogretmen PDF yukler,\nogrenci anlinda gorur",            GREEN),
    ("Cuma",               "Deneme sonuclari dijital\nolarak aktarilir",             AMBER),
    ("Cumartesi",          "AI, Ayse'yi riskli\nogrenci olarak isaretler",          RED),
    ("Sonraki\nHafta",     "Plan otomatik\nguncellenir",                            GREEN),
]

bar(s, 0.5, 3.62, 12.33, thick=0.055, color=ACCENT2)

step_w = 12.33 / len(steps)
sx = 0.5
for i, (day, desc, color) in enumerate(steps):
    cx2 = sx + step_w * i + step_w / 2 - 0.13
    dot(s, cx2, 3.49, d=0.26, color=color)
    tb(s, day, sx + step_w * i, 2.82, step_w, 0.65,
       size=13, bold=True, color=color, align=PP_ALIGN.CENTER)
    tb(s, desc, sx + step_w * i, 3.82, step_w, 0.9,
       size=13, align=PP_ALIGN.CENTER)

rect(s, 0.4, 4.88, 12.5, 0.72, fill=DARKGREEN)
tb(s, "->  Sistem, baskasından once fark etti.  Ve harekete gecti.",
   0.6, 4.97, 12.1, 0.55,
   size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)

details = [
    "Hicbir Excel tablosu olusturulmadi.  Hicbir ogretmenin onceden fark etmesi gerekmedi.",
    "Ogrenci verisi + kurum verisi  ->  otomatik olarak birlikte calisti.",
]
for i, d in enumerate(details):
    tb(s, d, 0.4, 5.78 + i * 0.46, 12.5, 0.42,
       size=15, color=GREY, align=PP_ALIGN.CENTER)

snum(s, 12)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 13 — REKABET AVANTAJLARI
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Rekabet Avantajlari")
bar(s, 0.4, 0.55, 2.5)

tb(s, "Neden StudyOS.  Neden simdi.",
   0.4, 0.65, 12.5, 0.85, size=40, bold=True)

advantages = [
    ("Uctan Uca Ekosistem",      "Ogrenci verisi ile kurum verisini baglayan tek platform",         ACCENT),
    ("Her Iki Tarafta AI",       "Kocluk . Planlama . Risk tespiti . Akademik danismanlik",         GREEN),
    ("Cift Yonlu Senkronizasyon","Tum veriler otomatik akar — manuel aktarim yok",                  AMBER),
    ("Optik Kagit Aktarimi",     "Kagit sinav sonuclari sureci degistirmeden dijitale tasinir",     ACCENT),
    ("Erken Uyari Sistemi",      "Riskli ogrenciler ogretmen farketmeden once tespit edilir",       RED),
    ("Veli Seffafligi",          "Yapisal aile gorunurlugu — guvenilir ve resmi",                  PURPLE),
    ("Gizlilik Oncelikli",       "KVKK ve GDPR uyumu basindan itibaren tasarimda mevcut",          GREEN),
    ("Olceklenebilir Yapi",      "Tek siniftan ulusal cok-subeli aglarına kadar esit etkinlik",    ACCENT2),
]

cell_w = 3.18
cell_h = 1.48
col_count = 4
mgn = 0.19
for i, (title, desc, color) in enumerate(advantages):
    col = i % col_count
    row = i // col_count
    cx3 = 0.4 + col * (cell_w + mgn)
    cy3 = 1.75 + row * (cell_h + 0.12)
    rrect(s, cx3, cy3, cell_w, cell_h, fill=BG_CARD)
    bar(s, cx3, cy3, cell_w, thick=0.06, color=color)
    tb(s, title, cx3 + 0.12, cy3 + 0.12, cell_w - 0.2, 0.5,
       size=14, bold=True, color=color)
    tb(s, desc,  cx3 + 0.12, cy3 + 0.62, cell_w - 0.2, 0.78,
       size=12, color=GREY)

snum(s, 13)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 14 — IS MODELI
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Is Modeli")
bar(s, 0.4, 0.55, 1.5)

tb(s, "Iki urun.  Coklu gelir akislari.",
   0.4, 0.65, 12.5, 0.85, size=38, bold=True)

# Ogrenci gelir karti
rrect(s, 0.4, 1.75, 5.85, 4.1, fill=BG_CARD)
bar(s, 0.4, 1.75, 5.85, thick=0.07, color=ACCENT)
tb(s, "Ogrenci Uygulamasi", 0.55, 1.88, 5.5, 0.55,
   size=19, bold=True, color=ACCENT)
student_rev = [
    ("Premium Abonelik",      "Aylik / yillik tam erisim"),
    ("Odulli Reklamlar",      "Uygulama ici odul kazanilir"),
    ("Uygulama Ici Satin Al", "Tek seferlik ozellik acma"),
]
ry = 2.55
for title, sub in student_rev:
    tb(s, title, 0.6, ry, 3.0, 0.42, size=16, bold=True)
    tb(s, sub, 3.75, ry + 0.02, 2.3, 0.4, size=14, color=GREY)
    ry += 0.85

# Kurum gelir karti
rrect(s, 6.9, 1.75, 6.03, 4.1, fill=BG_CARD)
bar(s, 6.9, 1.75, 6.03, thick=0.07, color=GREEN)
tb(s, "Kurumsal Platform", 7.05, 1.88, 5.7, 0.55,
   size=19, bold=True, color=GREEN)
inst_rev = [
    ("Kurum Lisansi",           "Kurum basina aylik / yillik"),
    ("Ogrenci Basina Lisans",   "Kayitli ogrenci basina ucret"),
    ("Premium AI Modulleri",    "Gelismis AI ozellikler"),
    ("Raporlama Paketleri",     "Gelismis analitik ve disari aktarim"),
]
ry = 2.55
for title, sub in inst_rev:
    tb(s, title, 7.0, ry, 3.3, 0.42, size=15, bold=True)
    tb(s, sub, 10.42, ry + 0.02, 2.4, 0.4, size=13, color=GREY)
    ry += 0.75

rect(s, 0.4, 6.05, 12.5, 0.75, fill=DARKGREEN)
tb(s, "Volan etkisi:  Kurum benimsemesi -> Daha fazla ogrenci -> Daha iyi AI -> Daha fazla kurum",
   0.6, 6.15, 12.1, 0.55,
   size=17, bold=True, color=GREEN, align=PP_ALIGN.CENTER)

snum(s, 14)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 15 — URUN YOL HARITASI
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Yol Haritasi")
bar(s, 0.4, 0.55, 1.5)

tb(s, "Dort asama.  Iki urun.  Tek yon.",
   0.4, 0.65, 12.5, 0.85, size=38, bold=True)

phases = [
    ("MVP",  ACCENT,
     "Giris . Calisma plani\nKonu takibi . Pomodoro\nIstatistik",
     "Kurum kurulumu . Siniflar\nOgretmen . Ogrenci\nPDF . Denemeler"),
    ("V2",   GREEN,
     "AI Koc\nDeneme analizi\nYanlis defteri\nAkilli tekrar",
     "AI analizi\nVeli paneli\nGelistirilmis raporlama"),
    ("V3",   AMBER,
     "Universite modulu\nPDF analizi . AI ozet\nFlashcard . Kariyer",
     "Canli ders\nMesajlasma\nOptik aktarim"),
    ("V4",   PURPLE,
     "Sosyal ozellikler\nCalisma gruplari\nWidget . Saat destegi",
     "Cok kurumlu yonetim\nLMS ozellikleri\nAPI entegrasyonlari"),
]
pw = 3.13
px = 0.4
for phase, color, stu, kur in phases:
    tag(s, phase, px, 1.75, pw - 0.08, 0.45, bg=color, fg=BG_DARK, size=16)
    rrect(s, px, 2.3, pw - 0.08, 2.55, fill=BG_CARD)
    bar(s, px, 2.3, pw - 0.08, thick=0.05, color=ACCENT)
    tb(s, "Ogrenci", px + 0.1, 2.35, pw - 0.25, 0.38,
       size=12, bold=True, color=ACCENT)
    tb(s, stu, px + 0.1, 2.72, pw - 0.25, 2.0, size=13)
    rrect(s, px, 5.0, pw - 0.08, 2.25, fill=BG_CARD)
    bar(s, px, 5.0, pw - 0.08, thick=0.05, color=GREEN)
    tb(s, "Kurum", px + 0.1, 5.05, pw - 0.25, 0.38,
       size=12, bold=True, color=GREEN)
    tb(s, kur, px + 0.1, 5.42, pw - 0.25, 1.7, size=13)
    px += pw + 0.1

snum(s, 15)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 16 — BEKLENEN ETKI
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Beklenen Etki")
bar(s, 0.4, 0.55, 1.8)

tb(s, "StudyOS hayata gecirildiginde neler degisir.",
   0.4, 0.65, 12.5, 0.85, size=36, bold=True)

impact_cols = [
    ("Ogrenciler Icin", ACCENT, [
        "Dagitik araclar -> tek sistem",
        "Kör calisma -> veri odakli hazirlik",
        "Izole emek -> AI rehberliginde ilerleme",
        "Tutarlilik otomatik takip edilir",
    ]),
    ("Kurumlar Icin", GREEN, [
        "Reaktif -> proaktif yonetim",
        "Manuel giris -> otomatik aktarim",
        "Gayri resmi iletisim -> yapisal kayitlar",
        "Riskli ogrenciler onceden tespit edilir",
    ]),
    ("Ekosistem Icin", AMBER, [
        "Ogrenci <-> kurum verisi ilk kez baglandi",
        "Veliler yapisal gorunum kazandi",
        "Kullanim arttikca AI gelistirilir",
        "Tum taraflar icin tek operasyonel katman",
    ]),
]
cx4 = 0.4
cw4 = 4.1
for heading, color, items in impact_cols:
    rrect(s, cx4, 1.75, cw4, 4.55, fill=BG_CARD)
    bar(s, cx4, 1.75, cw4, thick=0.07, color=color)
    tb(s, heading, cx4 + 0.12, 1.88, cw4 - 0.2, 0.55,
       size=19, bold=True, color=color)
    iy = 2.55
    for item in items:
        tb(s, "->  " + item, cx4 + 0.12, iy, cw4 - 0.2, 0.62, size=15)
        iy += 0.65
    cx4 += cw4 + 0.2

metrics = ["Aylik Aktif Kullanici", "Calisma Tutarliligi",
           "AI Etkilesim Orani", "Kurum Yenileme Orani",
           "Uyari Yanit Orani", "Memnuniyet Puani"]
mx = 0.4
for m in metrics:
    chip_w = len(m) * 0.115 + 0.5
    tag(s, m, mx, 6.6, chip_w, 0.52, bg=DARKBLUE, fg=ACCENT, size=12)
    mx += chip_w + 0.15

snum(s, 16)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 17 — UZUN VADELI VIZYON
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Uzun Vadeli Vizyon  .  3–5 Yil")
bar(s, 0.4, 0.55, 2.5)

tb(s, "Egitim isletim sistemi.",
   0.4, 0.65, 12.5, 0.85, size=42, bold=True)

visions = [
    ("Her ogrenci icin",
     "Sinav hazirligi'ndan universite mezuniyetine kadar tam akademik hayat tek uygulamada.\n"
     "Telefon  .  Akilli saat  .  Widget  .  Calisma gruplari.",
     ACCENT),
    ("Her kurum icin",
     "Tek sube veya yuzlerce sube — ayni platform, uyarlanan gorunum.\n"
     "Canli ders  .  LMS icerik  .  Acik API entegrasyonlari.",
     GREEN),
    ("Platform icin",
     "Ucuncu taraf icerik saglayicilari ve gelisitirciler acik API uzerinde insaat yapar.\n"
     "StudyOS, baskalarin da uzerine insaat yaptigi bir platforma donusur.",
     AMBER),
]
vy = 1.75
for heading, body, color in visions:
    rrect(s, 0.4, vy, 12.5, 1.55, fill=BG_CARD)
    bar(s, 0.4, vy, 0.1, thick=1.55, color=color)
    tb(s, heading, 0.72, vy + 0.12, 3.5, 0.55,
       size=19, bold=True, color=color)
    tb(s, body, 4.3, vy + 0.1, 8.4, 1.32, size=16)
    vy += 1.7

bar(s, 1.5, 6.8, 10.33, color=ACCENT)
tb(s, "Iki bagimsiz urun.  Tek ortak omurga.  Tek egitim isletim sistemi.",
   0, 6.92, 13.33, 0.5,
   size=18, bold=True, align=PP_ALIGN.CENTER)

snum(s, 17)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 18 — OZET
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
section(s, "Ozet")
bar(s, 0.4, 0.55, 1.2)

tb(s, "StudyOS bes maddede.",
   0.4, 0.65, 12.5, 0.8, size=42, bold=True)

points = [
    ("01", "Gercek Sorun",        "Ogrenciler ve kurumlar birbirinden kopuk.  En onemli veri gorunmez.",          ACCENT),
    ("02", "Gercek Cozum",        "Ortak bulut ile baglanan iki AI destekli urun otomatik veri paylasir.",        GREEN),
    ("03", "Gercek Farklilik",    "Kisisel calisma verisi ile kurum yonetimini baglayan tek platform.",           AMBER),
    ("04", "Gercek Is Modeli",    "Coklu gelir akislari ve her iki tarafin birbirini buyuttugu volan dinamigi.",  PURPLE),
    ("05", "Gercek Vizyon",       "Ogrenciden kuruma, bireyden ekosisteme egitim isletim sistemi.",               RGBColor(0xFB, 0xBF, 0x24)),
]

py = 1.65
ph = 0.93
for num, title, desc, color in points:
    rrect(s, 0.4, py, 12.5, ph - 0.06, fill=BG_CARD)
    bar(s, 0.4, py, 0.09, thick=ph - 0.06, color=color)
    tb(s, num,   0.65, py + 0.2, 0.7, 0.52,
       size=22, bold=True, color=color, align=PP_ALIGN.CENTER)
    tb(s, title, 1.45, py + 0.1, 3.0, 0.55, size=18, bold=True)
    tb(s, desc,  4.55, py + 0.13, 8.1, 0.65, size=15, color=GREY)
    py += ph + 0.04

bar(s, 2.0, 6.62, 9.33, color=ACCENT)
tb(s, "Iki urun.  Tek ekosistem.  Tek hedef.",
   0, 6.72, 13.33, 0.55,
   size=20, bold=True, align=PP_ALIGN.CENTER)

snum(s, 18)


# ══════════════════════════════════════════════════════════════════════════════
# SLAYT 19 — SORULAR VE TESEKURLER
# ══════════════════════════════════════════════════════════════════════════════
s = new_slide()
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG_DARK2

tb(s, "Tesekkurler.", 0, 1.4, 13.33, 1.4,
   size=72, bold=True, align=PP_ALIGN.CENTER)

bar(s, 3.5, 3.1, 6.33, color=ACCENT)

tb(s, "StudyOS", 0, 3.3, 13.33, 0.85,
   size=36, bold=True, align=PP_ALIGN.CENTER)
tb(s, "Yapay Zeka Destekli Egitim Ekosistemi",
   0, 4.12, 13.33, 0.55,
   size=20, color=GREY, align=PP_ALIGN.CENTER)

bar(s, 3.5, 4.85, 6.33, color=RGBColor(0x1E, 0x3A, 0x5F))

tb(s, "Sorularinizi bekliyoruz.",
   0, 5.05, 13.33, 0.62,
   size=22, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

tb(s, "docs/requirements/   .   docs/presentations/",
   0, 6.9, 13.33, 0.45,
   size=13, color=RGBColor(0x3A, 0x4A, 0x5E), align=PP_ALIGN.CENTER)

snum(s, 19)


# ── Kaydet ────────────────────────────────────────────────────────────────────
prs.save(OUT)
print("SAVED")
print("PATH=" + OUT)
print("SLIDES=" + str(len(prs.slides)))
